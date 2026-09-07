"""Render a Markdown draft into the formats a thesis actually gets submitted in.

Format policy, deliberately:

* **md** and **tex** are pure Python with no dependencies. LaTeX is a text transform, and
  a thesis is exactly the document class TeX exists for.
* **docx** needs `python-docx`, imported lazily so nobody installs it to export Markdown.
* **pdf** is delegated to an external toolchain (`pandoc`, else `pdflatex` on the LaTeX we
  already generate). Bundling a Python PDF renderer would mean hand-rolling pagination,
  hyphenation and font handling to produce something markedly worse than either tool -
  the honest move is to use the right one when it is installed and say so when it is not.

Every unavailable path raises `RenderError` with the exact command to fix it.
"""

import re
import shutil
import subprocess
import tempfile
from pathlib import Path

FORMATS = ("md", "tex", "docx", "pdf")
EXTENSIONS = {"md": ".md", "tex": ".tex", "docx": ".docx", "pdf": ".pdf"}

_PDF_TIMEOUT_SECONDS = 120

# LaTeX's ten special characters. Order is a genuine trap here: the replacements for
# `\`, `~` and `^` themselves contain braces, so substituting backslash first means the
# later `{`/`}` rules mangle `\textbackslash{}` into `\textbackslash\{\}`. Sentinels are
# substituted last, once no further escaping runs.
_SENTINELS = {
    "\\": "\x00bslash\x00",
    "~": "\x00tilde\x00",
    "^": "\x00caret\x00",
}
_SENTINEL_EXPANSIONS = {
    "\x00bslash\x00": r"\textbackslash{}",
    "\x00tilde\x00": r"\textasciitilde{}",
    "\x00caret\x00": r"\textasciicircum{}",
}
_TEX_ESCAPES = [
    ("&", r"\&"), ("%", r"\%"), ("$", r"\$"), ("#", r"\#"),
    ("_", r"\_"), ("{", r"\{"), ("}", r"\}"),
]


class RenderError(RuntimeError):
    """A format could not be produced, with an actionable reason."""


def escape_tex(text: str) -> str:
    for char, sentinel in _SENTINELS.items():
        text = text.replace(char, sentinel)
    for char, replacement in _TEX_ESCAPES:
        text = text.replace(char, replacement)
    for sentinel, expansion in _SENTINEL_EXPANSIONS.items():
        text = text.replace(sentinel, expansion)
    return text


def _inline_tex(text: str) -> str:
    """Escape, then restore the inline Markdown emphasis we understand."""
    out = escape_tex(text)
    out = re.sub(r"\*\*(.+?)\*\*", r"\\textbf{\1}", out)
    out = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"\\emph{\1}", out)
    out = re.sub(r"`(.+?)`", r"\\texttt{\1}", out)
    return out


def markdown_to_latex(markdown: str, title: str = "", author: str = "") -> str:
    """A compilable standalone LaTeX document.

    Covers the subset this project's drafts actually contain: ATX headings, unordered and
    ordered lists, and inline bold/italic/code. Anything else passes through as escaped
    text rather than being silently mangled.
    """
    body: list[str] = []
    list_env: str | None = None

    def close_list() -> None:
        nonlocal list_env
        if list_env:
            body.append(f"\\end{{{list_env}}}")
            list_env = None

    for raw in markdown.splitlines():
        line = raw.rstrip()

        # A Markdown horizontal rule would otherwise pass through as a literal `---`,
        # which TeX renders as an em-dash floating on its own line.
        if re.fullmatch(r"\s*(-{3,}|\*{3,}|_{3,})\s*", line):
            close_list()
            body.append(r"\par\noindent\hrulefill\par")
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading:
            close_list()
            level = len(heading.group(1))
            command = {1: "section", 2: "subsection", 3: "subsubsection", 4: "paragraph"}[level]
            body.append(f"\\{command}{{{_inline_tex(heading.group(2))}}}")
            continue

        bullet = re.match(r"^[-*+]\s+(.*)$", line)
        numbered = re.match(r"^\d+[.)]\s+(.*)$", line)
        if bullet or numbered:
            wanted = "itemize" if bullet else "enumerate"
            if list_env != wanted:
                close_list()
                list_env = wanted
                body.append(f"\\begin{{{wanted}}}")
            body.append(f"  \\item {_inline_tex((bullet or numbered).group(1))}")
            continue

        close_list()
        body.append("" if not line else _inline_tex(line))

    close_list()

    preamble = [
        r"\documentclass[12pt,a4paper]{article}",
        r"\usepackage[utf8]{inputenc}",
        r"\usepackage[T1]{fontenc}",
        r"\usepackage{hyperref}",
        r"\usepackage{parskip}",
    ]
    if title:
        preamble.append(f"\\title{{{escape_tex(title)}}}")
    if author:
        preamble.append(f"\\author{{{escape_tex(author)}}}")

    document = ["\\begin{document}"]
    if title:
        document.append(r"\maketitle")
    document += body + ["\\end{document}"]

    return "\n".join(preamble + [""] + document) + "\n"


def markdown_to_docx(markdown: str, destination: Path, title: str = "") -> Path:
    """Write a .docx. Requires python-docx."""
    try:
        from docx import Document
    except ImportError as exc:
        raise RenderError(
            "DOCX export needs python-docx:\n    pip install python-docx"
        ) from exc

    document = Document()
    if title:
        document.add_heading(title, level=0)

    for raw in markdown.splitlines():
        line = raw.rstrip()
        heading = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading:
            document.add_heading(heading.group(2), level=len(heading.group(1)))
            continue
        bullet = re.match(r"^[-*+]\s+(.*)$", line)
        if bullet:
            document.add_paragraph(bullet.group(1), style="List Bullet")
            continue
        numbered = re.match(r"^\d+[.)]\s+(.*)$", line)
        if numbered:
            document.add_paragraph(numbered.group(1), style="List Number")
            continue
        if line:
            document.add_paragraph(line)

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(destination))
    return destination


def pdf_toolchain() -> str | None:
    """Which external converter is available: 'pandoc', 'pdflatex', or None."""
    for tool in ("pandoc", "pdflatex"):
        if shutil.which(tool):
            return tool
    return None


def markdown_to_pdf(markdown: str, destination: Path, title: str = "") -> Path:
    """Produce a PDF using whichever external toolchain is installed."""
    tool = pdf_toolchain()
    if tool is None:
        raise RenderError(
            "PDF export needs an external converter, none found on PATH.\n"
            "  Install pandoc:   https://pandoc.org/installing.html\n"
            "  or a TeX distribution providing pdflatex (MiKTeX, TeX Live).\n"
            "Meanwhile `export: <dir>` with format 'tex' gives you a compilable source file."
        )

    destination.parent.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        if tool == "pandoc":
            source = tmp_path / "draft.md"
            source.write_text(markdown, encoding="utf-8")
            command = ["pandoc", str(source), "-o", str(destination)]
        else:
            source = tmp_path / "draft.tex"
            source.write_text(markdown_to_latex(markdown, title), encoding="utf-8")
            command = ["pdflatex", "-interaction=nonstopmode", "-halt-on-error",
                       "-output-directory", str(tmp_path), str(source)]

        result = subprocess.run(
            command, capture_output=True, text=True,
            timeout=_PDF_TIMEOUT_SECONDS, stdin=subprocess.DEVNULL,
        )
        if result.returncode != 0:
            tail = (result.stdout or result.stderr or "")[-1500:]
            raise RenderError(f"{tool} failed (exit {result.returncode}):\n{tail}")

        if tool == "pdflatex":
            produced = tmp_path / "draft.pdf"
            if not produced.exists():
                raise RenderError("pdflatex reported success but produced no PDF.")
            destination.write_bytes(produced.read_bytes())

    return destination


def render(markdown: str, fmt: str, destination: Path, title: str = "", author: str = "") -> Path:
    """Write `markdown` to `destination` in `fmt`. Returns the path written."""
    fmt = fmt.lower().lstrip(".")
    if fmt not in FORMATS:
        raise RenderError(f"Unknown format {fmt!r}. Available: {', '.join(FORMATS)}.")

    if fmt == "docx":
        return markdown_to_docx(markdown, destination, title)
    if fmt == "pdf":
        return markdown_to_pdf(markdown, destination, title)

    destination.parent.mkdir(parents=True, exist_ok=True)
    content = markdown if fmt == "md" else markdown_to_latex(markdown, title, author)
    destination.write_text(content, encoding="utf-8", newline="\n")
    return destination
